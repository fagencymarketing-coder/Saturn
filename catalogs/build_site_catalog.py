# -*- coding: utf-8 -*-
"""Единый список товаров для страницы каталога (Claude Design) и импорта в Tilda.

Источники: catalog_tech.json (87 препаратов, уже без OYYO) и seeds_tech.csv
(25 семян). Порядок — по ходу сезона. Значения — через typography.py.
Культуры для фильтра вытаскиваются из «Культуры_и_фазы» по словам;
«зерновые» раскрываются в пшеницу и ячмень.
"""
import csv, io, json, os, re, sys
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from typography import polish, price as price_fmt

HERE = os.path.dirname(os.path.abspath(__file__))

# раздел по ходу сезона → подразделы в порядке показа
SEASON = [
    # культуру у семян показывает фильтр — подраздел по культуре его дублировал
    ('Семена', ['Гибриды', 'Сорта']),
    ('Обработка семян и протравители', ['Обработка семян', 'Протравители']),
    ('Гербициды', ['Гербициды']),
    ('Удобрения и листовые подкормки', ['Удобрения', 'Листовые подкормки']),
    ('Фунгициды', ['Фунгициды']),
    ('Инсектициды', ['Инсектициды']),
    ('Десиканты', ['Десиканты']),
    ('Адъюванты', ['Адъюванты и спецпрепараты']),
]
SZR = {'Гербициды', 'Фунгициды', 'Инсектициды', 'Протравители', 'Десиканты'}
FERT = {'Удобрения', 'Листовые подкормки', 'Обработка семян'}
CROPS = [('пшениц', 'Пшеница'), ('ячмен', 'Ячмень'), ('рапс', 'Рапс'), ('подсолнечник', 'Подсолнечник'),
         (r'\bсо(я|и|ю)\b', 'Соя'), ('кукуруз', 'Кукуруза'), ('картофел', 'Картофель'),
         ('свёкл|свекл', 'Свёкла'), ('горох', 'Горох'), (r'\bл[её]н', 'Лён'), ('овощ', 'Овощи')]
BRAND = {'Волски': 'Волски Биохим', 'СЗР / адъюванты': ''}


# препараты для любых полевых культур показываются под каждым фильтром
UNIVERSAL = re.compile(r'все[хм]?\b[^.]*культур|полевые культуры|под посев|\bпары\b|при недостатке|семян и клубней')
VEG = re.compile(r'огур|томат|зелень')


def crops_of(text):
    t = polish(text).lower()
    if UNIVERSAL.search(t):
        return ['Все культуры']
    out = [name for pat, name in CROPS if re.search(pat, t)]
    if VEG.search(t) and 'Овощи' not in out:
        out.append('Овощи')
    if 'зернов' in t:
        out += [c for c in ('Пшеница', 'Ячмень') if c not in out]
    order = [n for _, n in CROPS]
    return sorted(set(out), key=order.index)


def section_of(sub):
    for sec, subs in SEASON:
        if sub in subs:
            return sec
    raise KeyError(sub)


rows = []
for r in json.load(io.open(os.path.join(HERE, 'catalog_tech.json'), encoding='utf-8')):
    cat = r['Категория']
    brand = BRAND.get(r['Бренд'], r['Бренд'])
    if cat in SZR:
        sub_line = polish(r['Состав_ДВ'])
    elif cat in FERT:
        sub_line = ' · '.join(x for x in (brand, cat.lower()) if x)
    else:
        sub_line = polish(r['Описание'])
    pack = polish(r.get('Фасовка_прайс') or r.get('Фасовка_каталог'))
    rows.append({
        'Раздел': section_of(cat), 'Подраздел': 'Адъюванты' if cat.startswith('Адъюванты') else cat,
        'Название': r['Название'], 'Бренд': brand, 'Подпись в карточке': sub_line,
        'Культуры (фильтр)': ', '.join(crops_of(r.get('Культуры_и_фазы'))),
        'Цена': price_fmt(r['Цена_итог']), 'Фасовка': '' if pack == 'см. прайс' else pack,
        'Фото': os.path.basename(r.get('Файл_вырезанный') or ''), '_sub': cat,
    })

SEED_PHOTO = {'Новосёл CL': 'seed-novosel-cl.jpg', 'Ampir 10 Express': 'seed-ampir-10.jpg',
              'Ampir 25 Express': 'seed-ampir-25.jpg'}
# у каждой культуры несколько кадров — соседние карточки сортов не повторяют друг друга
CROP_PHOTOS = {'Пшеница озимая': 'pshenitsa', 'Пшеница яровая': 'pshenitsa', 'Ячмень яровой': 'yachmen',
               'Соя': 'soya', 'Горох посевной': 'goroh', 'Гречиха': 'grechiha', 'Картофель': 'kartofel'}
_crop_dir = os.path.join(os.path.dirname(HERE), 'assets', 'crop-photos', 'cards')
_used = {}


def crop_photo(crop):
    slug = CROP_PHOTOS.get(crop)
    if not slug:
        return ''
    files = sorted(f for f in os.listdir(_crop_dir) if f.startswith(f'seed-crop-{slug}-'))
    i = _used.get(slug, 0); _used[slug] = i + 1
    return files[i % len(files)] if files else ''
CROP_TAG = {'Пшеница озимая': 'Пшеница', 'Пшеница яровая': 'Пшеница', 'Ячмень яровой': 'Ячмень',
            'Рапс яровой': 'Рапс', 'Горох посевной': 'Горох'}
repro = {}
for s in csv.DictReader(io.open(os.path.join(HERE, 'seeds.csv'), encoding='utf-8-sig'), delimiter=';'):
    key = s['Сорт'].replace('Гибрид ', '').replace('Новосел', 'Новосёл')
    repro.setdefault(key, []).append(s['Репродукция'].replace('РС1', 'РС 1'))
for s in csv.DictReader(io.open(os.path.join(HERE, 'seeds_tech.csv'), encoding='utf-8-sig'), delimiter=';'):
    crop, name = s['Культура'], s['Название']
    hybrid = s['Тип'].startswith('гибрид')
    sub = 'Гибриды' if hybrid else 'Сорта'
    kind = 'гибрид F1' if hybrid else ', '.join(dict.fromkeys(repro.get(name, [])))
    rows.append({
        'Раздел': 'Семена', 'Подраздел': sub, 'Название': name, 'Бренд': s['Оригинатор'] if 'не указан' not in s['Оригинатор'] else '',
        'Подпись в карточке': f'{crop} · {kind}', 'Культуры (фильтр)': CROP_TAG.get(crop, crop),
        'Цена': 'Запросить цену', 'Фасовка': '',
        'Фото': SEED_PHOTO.get(name) or crop_photo(crop), '_sub': sub,
    })

order = [(sec, sub) for sec, subs in SEASON for sub in subs]
def key(r):
    sub = r['_sub']
    sec = r['Раздел']
    return order.index((sec, sub if sec != 'Адъюванты' else 'Адъюванты и спецпрепараты'))
rows.sort(key=key)
cols = ['№', 'Раздел', 'Подраздел', 'Название', 'Бренд', 'Подпись в карточке', 'Культуры (фильтр)', 'Цена', 'Фасовка', 'Фото']
for i, r in enumerate(rows, 1):
    r['№'] = i

with io.open(os.path.join(HERE, 'site_catalog.csv'), 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.DictWriter(f, fieldnames=cols, delimiter=';', extrasaction='ignore'); w.writeheader(); w.writerows(rows)

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Каталог'
ws.append(cols)
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='141210')
for r in rows:
    ws.append([r[c] for c in cols])
for col, w_ in zip('ABCDEFGHIJ', (5, 26, 22, 34, 18, 46, 30, 22, 18, 34)):
    ws.column_dimensions[col].width = w_
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(vertical='top', wrap_text=True)
ws.freeze_panes = 'A2'
wb.save(os.path.join(HERE, 'Saturn-каталог-для-сайта.xlsx'))

from collections import Counter
print('позиций:', len(rows), '· с фото:', sum(1 for r in rows if r['Фото']))
for sec, _ in SEASON:
    print('  %-32s %d' % (sec, sum(1 for r in rows if r['Раздел'] == sec)))
cc = Counter(c for r in rows for c in r['Культуры (фильтр)'].split(', ') if c)
print('культуры:', ', '.join('%s %d' % kv for kv in cc.most_common()))
print('без культуры:', [r['Название'] for r in rows if not r['Культуры (фильтр)']])
