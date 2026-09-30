# -*- coding: utf-8 -*-
"""Сводная таблица каталога: данные каждой позиции + откуда они + что спорно.

Для заказчика и для нас: одна строка = одна позиция сайта (112).
Источники: подтверждённая таблица заказчика (28.09.2026), прайс «Сатурна» 2026,
каталоги производителей, Госреестр и сайты оригинаторов (семена), поиск 30.09.2026.
"""
import csv, json, os, re
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
def rd(p): return list(csv.DictReader(open(os.path.join(HERE, p), encoding='utf-8-sig'), delimiter=';'))
def key(s): return re.sub(r'[^a-zа-я0-9]', '', (s or '').lower().replace('ё', 'е').replace('k', 'к'))

site = rd('site_catalog.csv')
tech = {key(r['Название']): r for r in rd('catalog_tech.csv')}
seeds = {r['Название']: r for r in rd('seeds_tech.csv')}
seedprice = {r['Сорт']: r for r in rd('seeds.csv')}
units = {key(r['Название']): r for r in csv.DictReader(open(os.path.join(HERE, 'price_units.csv'), encoding='utf-8'), delimiter=';')}
reg = {}
for r in csv.DictReader(open(os.path.join(ROOT, 'assets/product-photos/реестр.csv'), encoding='utf-8-sig'), delimiter=';'):
    reg[key(r['товар'])] = r
ws = openpyxl.load_workbook(os.path.join(HERE, 'Saturn-сверка-трёх-источников.xlsx'))['Сверка трёх источников']
sv = {key(r[0]): r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}

SRC_LINK = {
    'Каталог АО «Фертика»': 'Каталог FERTIKA ProLine (PDF заказчика) · fertika.com',
    'Буклет ООО «Волски Биохим»': 'Буклет «Волски Биохим» (PDF заказчика) · volskybiochem.com',
    'Каталог ГК «Генезис» (2026)': 'Каталог ГК «Генезис» 2026 (PDF заказчика; бренд на сайте не показываем)',
    'Прайс ООО «Сатурн», сезон 2026 (стр. 2–5)': 'Прайс ООО «Сатурн», сезон 2026, стр. 2–5 (PDF заказчика)',
}
CONFIRMED = 'Таблица заказчика, подтверждена 28.09.2026'
UNVERIFIED = re.compile(r'seedculture|agronayka|direct\.farm|agroserver|10akrov|agromax|ainur|svoefermerstvo', re.I)

out = []
for r in site:
    n = r['Название']; k = key(n); t = tech.get(k); s = seeds.get(n)
    issues = []
    row = {'№': int(r['№']), 'Раздел': r['Раздел'], 'Подраздел': r['Подраздел'], 'Позиция': n,
           'Бренд / оригинатор': r['Бренд'], 'Цена на сайте': r['Цена'], 'Фасовка': r['Фасовка']}
    if s:
        sp = seedprice.get(n)
        row.update({'Описание': s['Описание'], 'Спелость': s['Спелость'], 'Вегетация': s['Вегетация'],
                    'Урожайность': s['Урожайность'], 'Качество': s['Качество']})
        row['Источник данных'] = s['Источник'] or 'КП ООО «Сатурн»'
        row['Источник цены'] = (f"{sp['Источник']}: {sp['Цена_руб_за_тонну_с_НДС']} ₽/т с НДС — на сайте не показываем, «Запросить цену»"
                                if sp else 'КП ООО «Сатурн» — цены нет, «Запросить цену»')
        if not s['Спелость'] or not s['Вегетация']:
            issues.append('Нет спелости/вегетации в открытых источниках' + (' — какой из Ampir 10/25 ранний, не ясно; поля скрыты' if n.startswith('Ampir') else ''))
        if UNVERIFIED.search(s['Источник'] or ''):
            issues.append('Часть цифр (белок/клейковина) — с сайта продавца семян, не оригинатора: непроверенный источник')
        if n == 'Алтайский 22':
            issues.append('Раньше было «пивоваренный»; по Госреестру — зернофуражный. Исправлено')
        if n == 'Юнион':
            issues.append('Год регистрации по Госреестру — 2022 (в прежних заметках 2024)')
        if 'не указан' in (s['Оригинатор'] or ''):
            issues.append('Оригинатор не указан в КП')
    elif t:
        row.update({'Описание': t['Описание'], 'Состав / ДВ': t['Состав_ДВ'], 'Норма расхода': t['Норма_расхода'],
                    'Культуры и фазы': t['Культуры_и_фазы'], 'Физ.-хим. показатели': t['Физпоказатели']})
        src = t['Источник'] or ''
        link = SRC_LINK.get(src, src)
        pg = reg.get(key(n.replace('FERTIKA ', ''))) or reg.get(k)
        if pg and pg.get('страница_товара') and not UNVERIFIED.search(pg['страница_товара']):
            link += ' · ' + pg['страница_товара']
        row['Источник данных'] = f'{CONFIRMED}; первоисточник — {link}' if src else CONFIRMED
        row['Источник цены'] = CONFIRMED + (' (цены удобрений сверены с прайсом «Сатурна» 2026)' if r['Раздел'] == 'Удобрения и листовые подкормки' else ' (СЗР — с прайса «Сатурна» 2026, стр. 2–5)')
        d = sv.get(k)
        if d:
            if d[4]: issues.append(f'Цена отличается от прайса «Сатурна» ({d[2]}); на сайте — по таблице заказчика ({d[3]})')
            if d[8]: issues.append(f'Норма расхода в прайсе «Сатурна» другая: «{d[6]}»; на сайте — по каталогу производителя')
            if d[11]: issues.append(f'Состав в прайсе «Сатурна»: «{d[10]}»')
        if t.get('Сверить'): issues.append('В таблице заказчика помечено «сверить»: ' + t['Сверить'])
        if re.search(r'(\d+(?:,\d+)?)\s*-\s*\1(?!\d)', t['Норма_расхода'] or ''):
            issues.append('В таблице норма «X–X» (одно число дважды) — показываем одно значение')
        if n == 'Волски Аминатор':
            issues.append('Норма «0,5» без единиц; официальных данных нет (продукт на регистрации)')
        if n == 'Микромак':
            issues.append('В таблице опечатка — иероглиф в тексте («всех基»); на сайте убран')
    else:
        row['Источник данных'] = '—'
    # единица цены
    u = units.get(k)
    if u and re.search(r'\d', r['Цена']):
        row['Единица цены — основание'] = u['Основание']
        issues.append(f'Единица «₽/{u["Единица"]}» не подтверждена заказчиком ({u["Основание"].split(":")[0] if "поиск" in u["Основание"] else "по фасовке"})')
    elif re.search(r'\d', r['Цена']) and r['Раздел'] != 'Семена':
        row['Единица цены — основание'] = 'нет данных — без единицы'
        issues.append('За что цена (л / кг / упаковка) — неизвестно')
    # фото
    if not r['Фото']:
        row['Фото — источник'] = 'нет фото'
        issues.append('Нет фото: на сайтах производителя и дилеров не найдено')
    elif r['Фото'].startswith('seed-'):
        row['Фото — источник'] = 'Фото культуры (Wikimedia Commons, CC0/PD — assets/crop-photos/README.md)' if 'crop' in r['Фото'] else 'Фото из КП / материалов заказчика'
    else:
        pg = reg.get(key(n.replace('FERTIKA ', ''))) or reg.get(k)
        row['Фото — источник'] = (pg.get('страница_товара') or pg.get('источник_изображения') or '') if pg else 'материалы заказчика'
        if pg and UNVERIFIED.search((pg.get('страница_товара') or '') + (pg.get('источник_изображения') or '')):
            issues.append('Фото с сайта продавца, а не производителя' + (' (банка 500 г, в прайсе фасовка 0,1 кг)' if 'Тифилагро' in n else ''))
    row['Спорно / уточнить'] = '\n'.join(f'• {x}' for x in issues)
    out.append(row)

COLS = ['№', 'Раздел', 'Подраздел', 'Позиция', 'Бренд / оригинатор', 'Цена на сайте', 'Единица цены — основание', 'Источник цены',
        'Фасовка', 'Описание', 'Состав / ДВ', 'Норма расхода', 'Культуры и фазы', 'Физ.-хим. показатели',
        'Спелость', 'Вегетация', 'Урожайность', 'Качество', 'Источник данных', 'Фото — источник', 'Спорно / уточнить']
wb = openpyxl.Workbook(); w = wb.active; w.title = 'Каталог — источники'
w.append(COLS)
for r in out: w.append([r.get(c, '') for c in COLS])
hdr = PatternFill('solid', fgColor='141210'); pink = PatternFill('solid', fgColor='FCE4D6'); thin = Side(style='thin', color='E7E3DE')
for c in w[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical='center')
widths = {'№': 5, 'Раздел': 16, 'Подраздел': 16, 'Позиция': 28, 'Бренд / оригинатор': 18, 'Цена на сайте': 18, 'Единица цены — основание': 26,
          'Источник цены': 30, 'Фасовка': 14, 'Описание': 40, 'Состав / ДВ': 30, 'Норма расхода': 34, 'Культуры и фазы': 34,
          'Физ.-хим. показатели': 24, 'Спелость': 14, 'Вегетация': 14, 'Урожайность': 26, 'Качество': 26, 'Источник данных': 44,
          'Фото — источник': 30, 'Спорно / уточнить': 60}
for i, c in enumerate(COLS, 1):
    w.column_dimensions[openpyxl.utils.get_column_letter(i)].width = widths[c]
for row in w.iter_rows(min_row=2):
    flag = bool(row[-1].value)
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = Border(bottom=thin)
    if flag: row[-1].fill = pink
w.freeze_panes = 'E2'; w.auto_filter.ref = w.dimensions
L = wb.create_sheet('Как читать', 0)
for line in [
    ('Каталог сайта sssaturn.ru — данные и источники', ''),
    ('Одна строка — одна позиция сайта (112).', ''),
    ('', ''),
    ('Главный источник', 'Таблица заказчика «Каталог и прайс», подтверждена 28.09.2026. Всё на сайте совпадает с ней по сути; оформление (тире, ₽, единицы) — единое.'),
    ('Первоисточники', 'Каталоги производителей (Фертика, Волски Биохим, Генезис), прайс ООО «Сатурн» 2026, для семян — Госреестр (gossortrf.ru) и сайты оригинаторов.'),
    ('Столбец «Спорно / уточнить»', 'Розовым — есть вопрос: расхождение с прайсом «Сатурна», единица цены не подтверждена, непроверенный источник, нет фото и т. п.'),
    ('Расхождения с прайсом', 'Прайс «Сатурна» 2026 — старее таблицы заказчика. На сайте везде действует таблица заказчика; прайсовое значение показано для сравнения.'),
    ('Решения без заказчика', 'Единицы цен, скрытая спелость Ampir, «15 культур», протоколы — см. docs/PRELAUNCH-2026-09-30.md. Пересматриваются по ответу заказчика.'),
]:
    L.append(list(line))
L['A1'].font = Font(bold=True, size=14); L.column_dimensions['A'].width = 28; L.column_dimensions['B'].width = 110
for row in L.iter_rows():
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
p = os.path.join(HERE, 'Saturn-каталог-источники-и-вопросы.xlsx'); wb.save(p)
print('позиций:', len(out), '· с вопросами:', sum(1 for r in out if r['Спорно / уточнить']))
