# -*- coding: utf-8 -*-
import json, io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
rows=json.load(io.open('catalogs/catalog_tech.json',encoding='utf-8'))
COLS=["Название","Бренд","Категория","Цена_итог","Фасовка_прайс","Описание",
 "Состав_ДВ","Норма_расхода","Культуры_и_фазы","Фасовка_каталог","Физпоказатели","Источник","Сверить"]
W=[38,18,24,30,22,54,60,34,58,20,38,30,12]
wb=Workbook(); ws=wb.active; ws.title="Каталог"
head=PatternFill("solid",fgColor="1A1A1A"); hf=Font(color="FFFFFF",bold=True,size=11)
thin=Side(style="thin",color="D9D5D0"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
ws.append(COLS)
for c in ws[1]: c.fill=head;c.font=hf;c.alignment=Alignment(vertical="center",horizontal="center",wrap_text=True)
for i,w in enumerate(W,1): ws.column_dimensions[get_column_letter(i)].width=w
ws.row_dimensions[1].height=32
alt=PatternFill("solid",fgColor="F6F5F3"); warn=PatternFill("solid",fgColor="FFE8E0")
szr=PatternFill("solid",fgColor="FFF3EE")
for n,r in enumerate(rows,start=2):
    ws.append([r[c] for c in COLS])
    is_szr = r["Бренд"]=="СЗР / адъюванты"
    for c in ws[n]:
        c.alignment=Alignment(vertical="top",wrap_text=True); c.border=bd
        if is_szr: c.fill=szr
        elif n%2==0: c.fill=alt
    ws.cell(n,1).font=Font(bold=True); ws.cell(n,4).font=Font(bold=True,color="FF4200")
    if r["Сверить"]: ws.cell(n,13).fill=warn
    ws.row_dimensions[n].height=80
ws.freeze_panes="B2"; ws.auto_filter.ref=f"A1:{get_column_letter(len(COLS))}{len(rows)+1}"

ws2=wb.create_sheet("Статус")
from collections import Counter
cnt=Counter(r["Категория"] for r in rows)
data=[["✅ ПРОВЕРЕНО",""],
["Цены удобрений","Сверены с прайсом ООО «Сатурн» (сезон 2026, с 12.01.2026). Расхождений нет."],
["Цены СЗР и адъювантов","Прочитаны заново с исходного PDF постранично (рендер 400 dpi). Старый OCR сдвигал цены на строку — например, Клео стоил «1065» вместо 4050. Исправлено."],
["Дубли","Удалено 7 дублей. Было 63 строки → 56 позиций удобрений."],
["Раздел СЗР","Добавлен целиком: гербициды, десиканты, протравители, инсектициды, фунгициды, адъюванты. Плюс «Фертика Газонное. Осень»."],
["",""],
["СОСТАВ КАТАЛОГА",""]]
for k,v in cnt.most_common(): data.append([k,f"{v} позиций"])
data+=[["ИТОГО",f"{len(rows)} позиций"],["",""],
["🔴 ОСТАЛОСЬ СПРОСИТЬ У ЗАКАЗЧИКА",""],
["Нормы расхода OYYO","В каталоге Генезиса только составы, норм нет. 12 позиций."],
["Цены OYYO и Relict Макро","Нет ни в прайсе, ни в каталогах → «Цена по запросу». 14 позиций."],
["Позиции «по запросу» в СЗР","Ампир ВР, Торпеда ВДГ, Тиль Про КС, Атлант Супер, Вега АнтиПена, Вега Баланс, Вега Клей, Вега Эмульс — цена в прайсе не указана."],
["",""],
["📌 ДЛЯ САЙТА",""],
["Условия","Цитата из прайса: «В стоимость препаратов входит доставка, хранение, полное сопровождение по циклу от посевной до уборки урожая»."],
["Контакты","На сайте показываем только Алексея Нилова +7 (960) 953-48-88 — решение заказчика."],
]
for r in data: ws2.append(r)
ws2.column_dimensions["A"].width=34; ws2.column_dimensions["B"].width=112
for row in ws2.iter_rows():
    for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)
for i,r in enumerate(data,1):
    if r[0].startswith(("✅","🔴","📌","СОСТАВ","ИТОГО")): ws2.cell(i,1).font=Font(bold=True,size=12,color="FF4200")
wb.save('catalogs/Saturn-каталог-полный-техданные.xlsx')
print("позиций:",len(rows))
